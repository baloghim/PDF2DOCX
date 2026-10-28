import csv
import io
from pathlib import Path
import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image
from pdf2docx_app.core.converter import convert
from pdf2docx_app.core.ocr import parse_tsv
from pdf2docx_app.core.builder import append_page


def paragraph(doc, text):
    return next(p for p in doc.paragraphs if p.text == text)


def test_pdf_styles_alignment_indents_spacing_and_page_breaks(tmp_path):
    path = tmp_path/'formatted.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=595, height=842)
        page.insert_textbox(pymupdf.Rect(50,40,545,90), 'FORMATTED TITLE', fontsize=20, fontname='hebo', align=1, color=(0,.2,.6))
        page.insert_text((70,140), 'Body intro', fontsize=12, fontname='tiro')
        page.insert_text((90,175), 'Indented paragraph', fontsize=12, fontname='tibo')
        page.insert_text((70,210), 'Italic note', fontsize=10, fontname='heit')
        page.insert_textbox(pymupdf.Rect(350,250,545,280), 'Signature', align=2, fontsize=11)
        page = pdf.new_page(width=595, height=842)
        page.insert_text((70,80), 'Second page')
        pdf.save(path)
    doc = Document(convert(path).path)
    title = paragraph(doc, 'FORMATTED TITLE')
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs[0].font.size.pt == 20
    assert title.runs[0].bold
    assert str(title.runs[0].font.color.rgb) == '003399'
    assert paragraph(doc, 'Body intro').runs[0].font.name == 'Times New Roman'
    assert paragraph(doc, 'Body intro').runs[0].font.size.pt == 12
    assert paragraph(doc, 'Indented paragraph').paragraph_format.left_indent.pt >= 19
    assert paragraph(doc, 'Italic note').runs[0].italic
    assert paragraph(doc, 'Signature').alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraph(doc, 'Body intro').paragraph_format.space_before.pt > 20
    assert len(doc.sections) == 2
    assert abs(doc.sections[0].page_width.pt-595) < 1
    assert abs(doc.sections[0].page_height.pt-842) < 1


def test_word_lists_preserve_number_values_and_are_editable(tmp_path):
    path = tmp_path/'lists.pdf'
    with pymupdf.open() as pdf:
        p = pdf.new_page()
        p.insert_text((70,80), '3. First item\n4. Second item', fontsize=11)
        pdf.save(path)
    doc = Document(convert(path).path)
    first, second = paragraph(doc,'First item'), paragraph(doc,'Second item')
    num_id = first._p.pPr.numPr.numId.val
    assert second._p.pPr.numPr.numId.val == num_id
    num = next(n for n in doc.part.numbering_part.element.findall(qn('w:num')) if n.numId == num_id)
    abstract_id = num.abstractNumId.val
    abstract = next(n for n in doc.part.numbering_part.element.findall(qn('w:abstractNum')) if int(n.get(qn('w:abstractNumId'))) == abstract_id)
    assert abstract.find('.//'+qn('w:start')).get(qn('w:val')) == '3'
    assert abstract.find('.//'+qn('w:lvlText')).get(qn('w:val')) == '%1.'
    assert first.paragraph_format.first_line_indent.pt < 0


def test_tables_and_images_keep_source_order_and_image_size(tmp_path):
    path = tmp_path/'order.pdf'
    image = Image.new('RGB', (100,40), 'red')
    buffer = io.BytesIO()
    image.save(buffer, format='PNG')
    with pymupdf.open() as pdf:
        p = pdf.new_page()
        p.insert_image(pymupdf.Rect(70,20,170,60), stream=buffer.getvalue())
        p.insert_text((70,90), 'Before table')
        for x in (70,200,350):
            p.draw_line((x,110),(x,190))
        for y in (110,150,190):
            p.draw_line((70,y),(350,y))
        for x,y,text in [(80,135,'A'),(210,135,'B'),(80,175,'C'),(210,175,'D')]:
            p.insert_text((x,y),text)
        p.insert_text((70,240),'After table')
        pdf.save(path)
    doc = Document(convert(path).path)
    assert len(doc.inline_shapes) == 1
    assert abs(doc.inline_shapes[0].width.pt-100) < 1
    blocks = list(doc.element.body)
    before = next(i for i,b in enumerate(blocks) if 'Before table' in ''.join(b.itertext()))
    table = next(i for i,b in enumerate(blocks) if b.tag == qn('w:tbl'))
    after = next(i for i,b in enumerate(blocks) if 'After table' in ''.join(b.itertext()))
    assert before < table < after
    assert abs(sum(c.width.pt for c in doc.tables[0].columns)-280) < 2


def test_ocr_geometry_retains_centered_title_and_indentation():
    fields = ['level','page_num','block_num','par_num','line_num','word_num','left','top','width','height','conf','text']
    stream = io.StringIO()
    writer = csv.DictWriter(stream,fieldnames=fields,delimiter='\t')
    writer.writeheader()
    for block, x,y,w,h,text in [(1,400,100,400,60,'TITLE'),(2,140,300,900,30,'Body text'),(3,220,420,500,30,'Indented text')]:
        writer.writerow(dict(zip(fields,[5,1,block,1,1,1,x,y,w,h,95,text])))
    model = parse_tsv(stream.getvalue(),600,840,2)
    doc = Document()
    append_page(doc,model,first=True)
    title = paragraph(doc,'TITLE')
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs[0].font.size.pt > paragraph(doc,'Body text').runs[0].font.size.pt
    assert paragraph(doc,'Indented text').paragraph_format.left_indent.pt >= 39
    assert paragraph(doc,'Indented text').paragraph_format.space_before.pt > 0


def test_short_address_lines_are_kept_without_text_boxes(tmp_path):
    path = tmp_path/'address.pdf'
    with pymupdf.open() as pdf:
        p = pdf.new_page()
        p.insert_text((70,80),'Recipient\nTown name\nStreet address', fontsize=11)
        pdf.save(path)
    doc = Document(convert(path).path)
    assert any(p.text == 'Recipient\nTown name\nStreet address' for p in doc.paragraphs)
    assert not doc.element.xpath('.//w:txbxContent')
