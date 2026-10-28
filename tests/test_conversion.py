import hashlib
import os
from pathlib import Path
import threading
import tkinter
import zipfile
import pytest
import pymupdf
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pdf2docx_app.core.converter import convert, validate_input
from pdf2docx_app.core.model import ConversionError, Cancelled
from pdf2docx_app.core.ocr import Tesseract

TEXT = 'Árvíztűrő tükörfúrógép. Ügyiratszám: AB-123/2026.'
FONT = os.environ.get('PDF2DOCX_TEST_FONT', str(Path(__file__).parent / 'assets' / 'DejaVuSans.ttf'))


def make_pdf(path, texts=(TEXT,)):
    pdf = pymupdf.open()
    for text in texts:
        page = pdf.new_page()
        page.insert_font(fontname='test', fontfile=FONT)
        page.insert_text((50, 70), text, fontname='test', fontsize=12)
    pdf.save(path)
    pdf.close()
    return path


def contents(path):
    doc = Document(path)
    return '\n'.join(p.text for p in doc.paragraphs)


def test_digital_unicode_multipage_and_source(tmp_path):
    source = make_pdf(tmp_path/'Árvíz (teszt).v2.PDF', (TEXT, '1. Első pont', '2. Második pont'))
    before = hashlib.sha256(source.read_bytes()).digest()
    class NoOCR:
        def recognize(self, *args):
            pytest.fail('Digital text must not use OCR')
    result = convert(source, ocr=NoOCR())
    text = contents(result.path)
    assert TEXT in text
    assert text.index('1. Első pont') < text.index('2. Második pont')
    assert result.pages == 3 and result.ocr_pages == 0
    assert hashlib.sha256(source.read_bytes()).digest() == before
    with zipfile.ZipFile(result.path) as docx:
        assert b'<w:t>' in docx.read('word/document.xml')
        assert not any(n.startswith('word/media/') for n in docx.namelist())


def test_input_errors(tmp_path):
    for paths in ([], ['a', 'b'], [tmp_path]):
        with pytest.raises(ConversionError):
            validate_input(paths)
    bad = tmp_path/'bad.pdf'
    bad.write_bytes(b'not a PDF')
    with pytest.raises(ConversionError, match='sérült'):
        convert(bad)
    txt = tmp_path/'file.txt'
    txt.write_text('test')
    with pytest.raises(ConversionError, match='fájltípus'):
        validate_input([txt])


def test_encrypted(tmp_path):
    pdf = pymupdf.open()
    pdf.new_page()
    source = tmp_path/'secret.pdf'
    pdf.save(source, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw='secret', owner_pw='owner')
    pdf.close()
    with pytest.raises(ConversionError, match='jelszóval'):
        convert(source)


def test_overwrite_and_failure_preserve_old(tmp_path, monkeypatch):
    source = make_pdf(tmp_path/'a.pdf')
    target = source.with_suffix('.docx')
    target.write_bytes(b'old data')
    with pytest.raises(ConversionError, match='létezik'):
        convert(source)
    def fail(*args):
        raise PermissionError()
    monkeypatch.setattr('pdf2docx_app.core.converter.os.replace', fail)
    with pytest.raises(ConversionError, match='menthető'):
        convert(source, overwrite=True)
    assert target.read_bytes() == b'old data'
    assert not list(tmp_path.glob('.pdf2docx-*'))


def test_cancel_no_partial(tmp_path):
    source = make_pdf(tmp_path/'a.pdf')
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(Cancelled):
        convert(source, cancel=cancel)
    assert not source.with_suffix('.docx').exists()


def test_target_race_no_clobber(tmp_path):
    source = make_pdf(tmp_path/'a.pdf')
    target = source.with_suffix('.docx')
    def progress(stage, *_):
        if stage == 'Dokumentum mentése…':
            target.write_bytes(b'created meanwhile')
    with pytest.raises(ConversionError):
        convert(source, progress=progress)
    assert target.read_bytes() == b'created meanwhile'


def test_tcl_paths():
    tcl = tkinter.Tcl()
    assert tcl.splitlist('{C:/Magyar ékezet/levél (2).pdf}') == ('C:/Magyar ékezet/levél (2).pdf',)
    assert len(tcl.splitlist('{C:/a b.pdf} C:/c.pdf')) == 2


def scan_page(pdf):
    image = Image.new('RGB', (1800, 700), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((80, 140), TEXT, fill='black', font=ImageFont.truetype(FONT, 48))
    import io
    stream = io.BytesIO()
    image.save(stream, format='PNG')
    page = pdf.new_page(width=600, height=233.33)
    page.insert_image(page.rect, stream=stream.getvalue())


@pytest.mark.parametrize('language', ['hun', 'hun+eng'])
def test_real_ocr_mixed_and_quality(tmp_path, language):
    source = make_pdf(tmp_path/'mixed.pdf', ('Digitális oldal',))
    with pymupdf.open(source) as pdf:
        scan_page(pdf)
        output = tmp_path/'scanned.pdf'
        pdf.save(output)
    result = convert(output, ocr=Tesseract(language))
    text = contents(result.path)
    assert result.pages == 2 and result.ocr_pages == 1
    assert 'Digitális oldal' in text
    assert TEXT in text  # exact synthetic reference: CER and WER are zero
    assert text.count('Árvíztűrő') == 1


def test_missing_ocr(tmp_path):
    source = tmp_path/'scan.pdf'
    with pymupdf.open() as pdf:
        scan_page(pdf)
        pdf.save(source)
    ocr = Tesseract()
    ocr.command = None
    with pytest.raises(ConversionError, match='OCR-motor'):
        convert(source, ocr=ocr)
    assert not source.with_suffix('.docx').exists()


def test_same_page_scan_and_text(tmp_path):
    source = tmp_path/'same-page.pdf'
    with pymupdf.open() as pdf:
        scan_page(pdf)
        page = pdf[0]
        page.set_mediabox(pymupdf.Rect(0, 0, 600, 500))
        page.insert_text((50, 50), 'Digital footer')
        pdf.save(source)
    result = convert(source)
    text = contents(result.path)
    assert TEXT in text and 'Digital footer' in text
    assert text.count('Árvíztűrő') == 1


def test_missing_language(tmp_path):
    source = tmp_path/'scan.pdf'
    with pymupdf.open() as pdf:
        scan_page(pdf)
        pdf.save(source)
    with pytest.raises(ConversionError, match='nyelvi'):
        convert(source, ocr=Tesseract('nonexistent'))


def test_paragraph_join_and_bold(tmp_path):
    source = tmp_path/'paragraph.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((50, 60), 'First line\ncontinued line', fontname='hebo', fontsize=11)
        pdf.save(source)
    result = convert(source)
    doc = Document(result.path)
    assert any(p.text == 'First line continued line' for p in doc.paragraphs)
    assert any(r.bold for p in doc.paragraphs for r in p.runs)


def test_actual_table(tmp_path):
    source = tmp_path/'table.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        for x in (50, 200, 350):
            page.draw_line((x, 100), (x, 180))
        for y in (100, 140, 180):
            page.draw_line((50, y), (350, y))
        for x, y, text in [(60, 125, 'A'), (210, 125, 'B'), (60, 165, 'C'), (210, 165, 'D')]:
            page.insert_text((x, y), text)
        pdf.save(source)
    result = convert(source)
    doc = Document(result.path)
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [['A', 'B'], ['C', 'D']]
    assert not any(p.text.strip() for p in doc.paragraphs)
