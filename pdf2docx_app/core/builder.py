"""Create ordinary Word paragraphs, lists, tables and inline images with source styles."""
import io
import re
from dataclasses import replace
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
MARKER = re.compile(r'^\s*(?:(?P<number>\d{1,3})(?P<suffix>[.)])|(?P<bullet>[•●▪]))\s+')


def numbering(doc, marker, previous):
    number = int(marker['number']) if marker['number'] else None
    suffix = marker['suffix'] or '•'
    if previous and previous[1] == suffix and (number is None or number == previous[2]+1):
        return previous[0], (previous[0], suffix, number)
    root = doc.part.numbering_part.element
    ids = [int(e.get(qn('w:abstractNumId'))) for e in root.findall(qn('w:abstractNum'))]
    abstract_id = max(ids, default=-1)+1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    for tag, value in [('start', str(number if number is not None else 1)),
                       ('numFmt', 'decimal' if number is not None else 'bullet'),
                       ('suff', 'space'),
                       ('lvlText', '%1'+suffix if number is not None else '•')]:
        child = OxmlElement('w:'+tag)
        child.set(qn('w:val'), value)
        level.append(child)
    abstract.append(level)
    root.insert_element_before(abstract, "w:num", "w:numIdMacAtCleanup")
    num = root.add_num(abstract_id)
    return num.numId, (num.numId, suffix, number)


def trim_marker(runs, count):
    result = []
    for run in runs:
        if count >= len(run.text):
            count -= len(run.text)
        else:
            result.append(replace(run, text=run.text[count:]))
            count = 0
    return result


def append_page(doc, model, first=False):
    boxes = [p.bbox for p in model.paragraphs if p.bbox] + model.table_boxes + model.image_boxes
    left = min((b[0] for b in boxes), default=56.7)
    right = max((b[2] for b in boxes), default=model.width-56.7)
    top = min((b[1] for b in boxes), default=56.7)
    bottom = max((b[3] for b in boxes), default=model.height-56.7)
    margin_left = max(18, min(90, left))
    margin_right = max(18, min(90, model.width-right))
    margin_top = max(18, min(90, top))
    margin_bottom = max(18, min(72, model.height-bottom))
    section = doc.sections[0] if first else doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width, section.page_height = Pt(model.width), Pt(model.height)
    section.left_margin, section.right_margin = Pt(margin_left), Pt(margin_right)
    section.top_margin, section.bottom_margin = Pt(margin_top), Pt(margin_bottom)
    content_width = model.width-margin_left-margin_right
    elements = []
    for p in model.paragraphs:
        elements.append(('paragraph', p, p.bbox))
    for index, rows in enumerate(model.tables):
        box = model.table_boxes[index] if index < len(model.table_boxes) else None
        elements.append(('table', (index, rows), box))
    for index, image in enumerate(model.images):
        box = model.image_boxes[index] if index < len(model.image_boxes) else None
        elements.append(('image', image, box))
    elements.sort(key=lambda item: (item[2][1], item[2][0]) if item[2] else (float('inf'), 0))
    previous_bottom = margin_top
    previous_list = None
    for kind, value, box in elements:
        gap = max(0, min(72, box[1]-previous_bottom)) if box else 0
        if kind == 'paragraph':
            paragraph = value
            p = doc.add_paragraph(style='Heading 1' if paragraph.heading else None)
            fmt = p.paragraph_format
            fmt.space_before, fmt.space_after = Pt(gap), Pt(0)
            fmt.line_spacing = Pt(paragraph.line_height) if paragraph.line_height else paragraph.line_spacing
            fmt.keep_with_next = paragraph.heading
            fmt.widow_control = True
            p.alignment = ALIGN.get(paragraph.alignment, WD_ALIGN_PARAGRAPH.LEFT)
            if box and paragraph.alignment in ('left', 'justify'):
                fmt.left_indent = Pt(max(0, (paragraph.body_left if paragraph.body_left is not None else box[0])-margin_left))
                fmt.first_line_indent = Pt(paragraph.first_line_indent)
                if paragraph.right_edge is not None:
                    fmt.right_indent = Pt(max(0, model.width-margin_right-paragraph.right_edge))
            runs = paragraph.runs
            marker = MARKER.match(paragraph.text) if paragraph.list_item else None
            if marker:
                num_id, previous_list = numbering(doc, marker, previous_list)
                num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_ilvl().val = 0
                num_pr.get_or_add_numId().val = num_id
                fmt.left_indent = Pt(max(14, (box[0]-margin_left+14) if box else 14))
                fmt.first_line_indent = Pt(-14)
                runs = trim_marker(runs, marker.end())
            else:
                previous_list = None
            for span in runs:
                r = p.add_run(span.text)
                r.bold, r.italic = span.bold, span.italic
                if span.size:
                    r.font.size = Pt(max(6, min(72, span.size)))
                if span.font:
                    r.font.name = span.font
                r.font.color.rgb = RGBColor.from_string(f'{span.color & 0xffffff:06X}')
                r.font.superscript = span.superscript
        elif kind == 'table':
            previous_list = None
            index, rows = value
            table = doc.add_table(rows=0, cols=max(map(len, rows)))
            table.style = 'Table Grid'
            widths = model.table_columns[index] if index < len(model.table_columns) else []
            if not widths or len(widths) != len(table.columns) or not all(widths):
                total = box[2]-box[0] if box else content_width
                widths = [total/len(table.columns)] * len(table.columns)
            scale = min(1, content_width/sum(widths))
            table.autofit = False
            table_width = table._tbl.tblPr.find(qn('w:tblW'))
            table_width.set(qn('w:type'), 'dxa')
            table_width.set(qn('w:w'), str(round(sum(widths)*scale*20)))
            if gap and doc.paragraphs and list(doc.element.body)[-2].tag == qn('w:p'):
                doc.paragraphs[-1].paragraph_format.space_after = Pt(gap)
            for column, width in zip(table.columns, widths):
                column.width = Pt(width*scale)
            if box:
                indent = OxmlElement('w:tblInd')
                indent.set(qn('w:w'), str(round(max(0, box[0]-margin_left)*20)))
                indent.set(qn('w:type'), 'dxa')
                table._tbl.tblPr.insert_element_before(indent, 'w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook')
            for row in rows:
                cells = table.add_row().cells
                for cell, text, width in zip(cells, row, widths):
                    cell.width = Pt(width*scale)
                    cell.text = text
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
        else:
            previous_list = None
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(gap)
            p.paragraph_format.space_after = Pt(0)
            if box:
                p.paragraph_format.left_indent = Pt(max(0, box[0]-margin_left))
            width = min(content_width, box[2]-box[0]) if box else 85
            p.add_run().add_picture(io.BytesIO(value), width=Pt(max(1, width)))
        if box:
            previous_bottom = max(previous_bottom, box[3])
