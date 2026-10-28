import os
from pathlib import Path
import tempfile
import threading
from dataclasses import dataclass
import pymupdf
from docx import Document
from docx.shared import Cm, Pt
from .extract import extract, usable
from .model import ConversionError, Cancelled
from .ocr import Tesseract
from .builder import append_page


@dataclass
class Result:
    path: Path
    pages: int
    ocr_pages: int
    uncertain_words: int


def validate_input(paths):
    if len(paths) != 1:
        raise ConversionError('Egyszerre csak egy PDF-fájl konvertálható.')
    path = Path(paths[0])
    if not path.is_file():
        raise ConversionError('A kiválasztott fájl nem létezik, vagy nem normál fájl.')
    if path.suffix.lower() != '.pdf':
        raise ConversionError('Hibás fájltípus! A program kizárólag PDF-dokumentumokat fogad.')
    return path


def open_pdf(path):
    try:
        pdf = pymupdf.open(path)
        if not pdf.is_pdf:
            pdf.close()
            raise ValueError()
    except Exception:
        raise ConversionError('A kiválasztott PDF-dokumentum nem nyitható meg, vagy sérült.') from None
    if pdf.needs_pass:
        pdf.close()
        raise ConversionError('A dokumentum jelszóval védett, ezért nem dolgozható fel.')
    if not len(pdf):
        pdf.close()
        raise ConversionError('A PDF-dokumentum üres.')
    return pdf


def convert(source, *, overwrite=False, progress=lambda *args: None, cancel=None, ocr=None):
    source = validate_input([source])
    target = source.with_suffix('.docx')
    cancel = cancel or threading.Event()
    ocr = ocr or Tesseract()
    if target.exists() and not overwrite:
        raise ConversionError('A Word-dokumentum már létezik; felülírási engedély szükséges.')
    temporary = None
    try:
        with open_pdf(source) as pdf:
            doc = Document()
            section = doc.sections[0]
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2)
            normal = doc.styles['Normal']
            normal.font.name, normal.font.size = 'Calibri', Pt(11)
            normal.paragraph_format.space_after = Pt(6)
            ocr_count = uncertain = 0
            for index, page in enumerate(pdf):
                if cancel.is_set():
                    raise Cancelled('A feldolgozás megszakadt.')
                progress('Dokumentum elemzése…', index, len(pdf))
                text = page.get_text()
                if usable(text):
                    progress('Szöveg kinyerése…', index, len(pdf))
                    model = extract(page)
                    # OCR separate large image regions only when no existing text
                    # overlaps them. Existing OCR layers are never duplicated.
                    words = page.get_text('words')
                    regions = []
                    for image in page.get_images(full=True):
                        for rect in page.get_image_rects(image[0]):
                            if rect.get_area() < page.rect.get_area() * .25:
                                continue
                            if any(rect.intersects(pymupdf.Rect(w[:4])) for w in words):
                                continue
                            if any(rect.intersects(previous) for previous in regions):
                                continue
                            regions.append(rect)
                    for rect in regions:
                        progress('Szövegfelismerés…', index, len(pdf))
                        addition = ocr.recognize(page, cancel, clip=rect)
                        model.paragraphs.extend(addition.paragraphs)
                        model.uncertain_words += addition.uncertain_words
                        model.ocr = True
                else:
                    progress('Szövegfelismerés…', index, len(pdf))
                    model = ocr.recognize(page, cancel)
                ocr_count += model.ocr
                uncertain += model.uncertain_words
                append_page(doc, model, first=(index == 0))
                progress('Word-dokumentum létrehozása…', index+1, len(pdf))
            if not any(p.text.strip() for p in doc.paragraphs) and not doc.tables:
                raise ConversionError('Nem sikerült szerkeszthető szöveget kinyerni a dokumentumból.')
            progress('Dokumentum mentése…', len(pdf), len(pdf))
            fd, temporary = tempfile.mkstemp(prefix='.pdf2docx-', suffix='.docx', dir=target.parent)
            os.close(fd)
            doc.save(temporary)
            verified = Document(temporary)
            if not any(p.text.strip() for p in verified.paragraphs) and not verified.tables:
                raise ConversionError('A létrehozott Word-dokumentum üres.')
            if cancel.is_set():
                raise Cancelled('A feldolgozás megszakadt.')
            if overwrite:
                os.replace(temporary, target)
            else:
                # Atomic no-clobber publication, even if a target appeared during conversion.
                os.link(temporary, target)
                os.unlink(temporary)
            temporary = None
            return Result(target, len(pdf), ocr_count, uncertain)
    except ConversionError:
        raise
    except OSError:
        raise ConversionError('A dokumentum nem menthető. Ellenőrizze az írási jogosultságot, a szabad tárhelyet, és zárja be a célfájlt a Wordben.') from None
    except Exception:
        raise ConversionError('Váratlan feldolgozási hiba történt. A dokumentum nem készült el.') from None
    finally:
        if temporary:
            Path(temporary).unlink(missing_ok=True)
